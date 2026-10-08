"""Read-only AICT inventory inspection and create-only, allowlisted review copies."""

import argparse
from collections import defaultdict
from contextlib import contextmanager
from datetime import datetime
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import re
import tempfile
from urllib.parse import unquote, urlsplit
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.packaging.core import DocumentProperties


MODEL_HEADER = "Model / Sub-component"
ID_HEADER = "AICT ID"
MAX_BYTES = 128 * 1024 * 1024


class WorkbookError(ValueError):
    """A workbook or sharing request cannot be processed safely."""


def _sha256(data):
    return hashlib.sha256(data).hexdigest()


def _text(value):
    return "" if value is None else str(value).strip()


@contextmanager
def _source(path):
    source = Path(path).resolve()
    if source.suffix.lower() != ".xlsx":
        raise WorkbookError("Only .xlsx inputs are supported; do not bypass protection.")
    try:
        if source.stat().st_size > MAX_BYTES:
            raise WorkbookError("Workbook exceeds the 128 MiB input safety limit.")
        data = source.read_bytes()
        if data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
            raise WorkbookError("Encrypted or legacy Office container; stop and recover a trusted copy.")
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if any(entry.flag_bits & 1 for entry in entries):
                raise WorkbookError("Encrypted ZIP workbook; protection must not be bypassed.")
            if sum(entry.file_size for entry in entries) > MAX_BYTES:
                raise WorkbookError("Expanded workbook exceeds the 128 MiB safety limit.")
            names = [entry.filename for entry in entries]
            if len(names) != len(set(names)):
                raise WorkbookError("Corrupt workbook: duplicate ZIP members.")
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise WorkbookError("Invalid XLSX package: required workbook parts are missing.")
            if archive.testzip() is not None:
                raise WorkbookError("Corrupt workbook ZIP checksum; stop without repair or revert.")
        book = load_workbook(BytesIO(data), data_only=False, keep_links=False)
    except WorkbookError:
        raise
    except BadZipFile as error:
        raise WorkbookError("Invalid ZIP or corrupt XLSX; stop without repair or revert.") from error
    except Exception as error:
        raise WorkbookError("Cannot read workbook: unreadable, encrypted, or corrupt XLSX; stop.") from error
    try:
        yield source, data, book, names
    finally:
        book.close()


def _inventory_sheet(book, sheet, first_n):
    if first_n < 1:
        raise WorkbookError("--first-n must be positive.")
    if sheet is not None and sheet not in book.sheetnames:
        raise WorkbookError("Requested sheet does not exist.")
    matches = []
    for candidate in ([book[sheet]] if sheet is not None else book.worksheets):
        for row in candidate.iter_rows(min_row=1, max_row=min(first_n, candidate.max_row)):
            headers = defaultdict(list)
            for cell in row:
                if cell.data_type != "f":
                    headers[_text(cell.value)].append(cell.column)
            if MODEL_HEADER in headers and ID_HEADER in headers:
                if any(name and len(columns) > 1 for name, columns in headers.items()):
                    raise WorkbookError("Ambiguous duplicate named headers; select an unambiguous sheet.")
                matches.append((candidate, row[0].row, headers))
    if len(matches) != 1:
        raise WorkbookError("Expected one Model / Sub-component and AICT ID header row within --first-n; specify --sheet or correct ambiguity.")
    candidate, header_row, headers = matches[0]
    return candidate, header_row, {name: columns[0] for name, columns in headers.items() if name}


def _model_rows(sheet, header_row, headers):
    return [row for row in range(header_row + 1, sheet.max_row + 1)
            if _text(sheet.cell(row, headers[MODEL_HEADER]).value)]


def _protected(book):
    return bool(book.security and (book.security.lockStructure or book.security.lockWindows
                or book.security.workbookPassword or book.security.revisionsPassword))


def _assert_unchanged(source, original):
    if _sha256(source.read_bytes()) != _sha256(original):
        raise WorkbookError("Source changed during processing; stop and inspect independently.")


def inspect_workbook(input_path, sheet=None, first_n=50):
    with _source(input_path) as (source, original, book, names):
        selected, header_row, headers = _inventory_sheet(book, sheet, first_n)
        rows = _model_rows(selected, header_row, headers)
        identities = defaultdict(list)
        missing = []
        for row in rows:
            value = _text(selected.cell(row, headers[ID_HEADER]).value)
            if not value or value.casefold() == "missing":
                missing.append(row)
            else:
                identities[value].append(row)
        hidden = sum(bool(selected.row_dimensions[row].hidden) for row in rows)
        _assert_unchanged(source, original)
        return {
            "source_sha256": _sha256(original), "sheet": selected.title,
            "header_row": header_row, "model_rows": rows, "model_row_count": len(rows),
            "missing_id_rows": missing,
            "duplicate_ids": {value: occurrences for value, occurrences in identities.items() if len(occurrences) > 1},
            "hidden_model_rows": hidden, "visible_model_rows": len(rows) - hidden,
            "retained_model_rows": len(rows),
            "sheets": [{"name": tab.title, "state": tab.sheet_state,
                        "protected": bool(tab.protection.sheet)} for tab in book.worksheets],
            "workbook_protected": _protected(book),
            "hidden_column_ranges": [key for key, dimension in selected.column_dimensions.items() if dimension.hidden],
            "formula_cells": sum(cell.data_type == "f" for row in selected for cell in row),
            "comment_cells": sum(cell.comment is not None for row in selected for cell in row),
            "external_link_parts": sum(name.startswith("xl/externalLinks/") for name in names),
            "macro_parts": sum("vba" in name.lower() for name in names),
            "assessment": "Generic risk assessment, not a complete OOXML schema or policy validator. Hidden content is not redacted.",
        }


def _sharing_config(config):
    if not isinstance(config, dict) or set(config) != {"columns", "trusted_hosts"}:
        raise WorkbookError("Columns config must contain exactly columns and trusted_hosts.")
    columns, hosts = config["columns"], config["trusted_hosts"]
    if not isinstance(columns, list) or not columns or any(not isinstance(name, str) or not name.strip() for name in columns):
        raise WorkbookError("columns must be a nonempty list of exact named headers.")
    if len(set(columns)) != len(columns):
        raise WorkbookError("columns must not contain duplicates.")
    if not isinstance(hosts, list) or any(not isinstance(host, str) or not re.fullmatch(r"[a-z0-9]+(?:[.-][a-z0-9]+)*", host) for host in hosts):
        raise WorkbookError("trusted_hosts must be a list of lowercase exact hostnames, without wildcards, ports, or URLs.")
    return columns, set(hosts)


def _approved_link(link, hosts):
    if link.location or not link.target:
        raise WorkbookError("Internal or ambiguous hyperlinks are not allowed in review exports.")
    target = link.target
    decoded = unquote(target)
    if any(character.isspace() or ord(character) < 32 or ord(character) == 127 for character in decoded) or "\\" in decoded:
        raise WorkbookError("Unsafe hyperlink rejected.")
    try:
        parsed = urlsplit(target)
        approved = (parsed.scheme == "https" and parsed.hostname in hosts
                    and parsed.username is None and parsed.password is None
                    and parsed.port in (None, 443) and not parsed.query and not parsed.fragment
                    and "?" not in decoded and "#" not in decoded)
    except ValueError as error:
        raise WorkbookError("Malformed hyperlink rejected.") from error
    if not approved:
        raise WorkbookError("Hyperlink rejected: require HTTPS, an exact trusted host, no credentials, query, or fragment.")
    return target


def _set_value(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    elif value is None:
        cell.data_type = "inlineStr"
        cell.number_format = "@"


def _validate_copy(path, columns, records):
    with _source(path) as (_, _, book, names):
        if book.sheetnames != ["Review"] or book.active.sheet_state != "visible":
            raise WorkbookError("Export validation failed: unexpected sheet structure.")
        review = book.active
        if review.max_row != len(records) + 1 or review.max_column != len(columns):
            raise WorkbookError("Export validation failed: row or column count mismatch.")
        if [cell.value for cell in review[1]] != columns:
            raise WorkbookError("Export validation failed: column allowlist mismatch.")
        for row_number, record in enumerate(records, 2):
            for column_number, (value, target) in enumerate(record, 1):
                cell = review.cell(row_number, column_number)
                if cell.value != value or type(cell.value) is not type(value):
                    if not (value is None and cell.value is None):
                        raise WorkbookError("Export validation failed: value mismatch.")
                if cell.data_type == "f" or cell.comment is not None or (cell.hyperlink.target if cell.hyperlink else None) != target:
                    raise WorkbookError("Export validation failed: formula, comment, or hyperlink mismatch.")
        if any(dimension.hidden for dimension in review.row_dimensions.values()) or any(dimension.hidden for dimension in review.column_dimensions.values()):
            raise WorkbookError("Export validation failed: hidden content.")
        if any("vba" in name.lower() or name.startswith(("xl/externalLinks/", "xl/comments", "xl/worksheets/sheet2", "docProps/custom")) for name in names):
            raise WorkbookError("Export validation failed: unexpected sensitive package parts.")
        if book.properties.creator or book.properties.lastModifiedBy or book.properties.description:
            raise WorkbookError("Export validation failed: identifying metadata.")


def sharing_export(input_path, output_path, config, sheet=None, first_n=50):
    columns, hosts = _sharing_config(config)
    output = Path(output_path).absolute()
    if output.suffix.lower() != ".xlsx":
        raise WorkbookError("Sharing output must be a new .xlsx file.")
    if output.resolve() == Path(input_path).resolve() or output.exists() or output.is_symlink():
        raise WorkbookError("Output must differ from input and must not already exist; no overwrite.")
    with _source(input_path) as (source, original, book, _):
        selected, header_row, headers = _inventory_sheet(book, sheet, first_n)
        if _protected(book) or any(tab.protection.sheet for tab in book.worksheets):
            raise WorkbookError("Protected workbook; stop without bypassing protection.")
        if any(column not in headers for column in columns):
            raise WorkbookError("Allowlisted column is absent from the detected header row.")
        rows = _model_rows(selected, header_row, headers)
        records = []
        for row in rows:
            record = []
            for column in columns:
                cell = selected.cell(row, headers[column])
                if cell.data_type == "f":
                    raise WorkbookError(f"Selected formula at row {row}; export rejected, do not use cached results.")
                value = cell.value
                if value == "":
                    value = None
                record.append((value, _approved_link(cell.hyperlink, hosts) if cell.hyperlink else None))
            records.append(record)
        exported = Workbook()
        review = exported.active
        review.title = "Review"
        fixed_date = datetime(2000, 1, 1)
        exported.properties = DocumentProperties(creator="", lastModifiedBy="", created=fixed_date, modified=fixed_date)
        for column_number, name in enumerate(columns, 1):
            _set_value(review.cell(1, column_number), name)
        for row_number, record in enumerate(records, 2):
            for column_number, (value, target) in enumerate(record, 1):
                destination = review.cell(row_number, column_number)
                _set_value(destination, value)
                if target:
                    destination.hyperlink = target
        temporary = None
        try:
            descriptor, temporary_name = tempfile.mkstemp(prefix=".aict-review-", suffix=".xlsx", dir=output.parent)
            temporary = Path(temporary_name)
            with os.fdopen(descriptor, "wb") as stream:
                exported.save(stream)
                stream.flush()
                os.fsync(stream.fileno())
            _validate_copy(temporary, columns, records)
            output_hash = _sha256(temporary.read_bytes())
            _assert_unchanged(source, original)
            os.link(temporary, output)
            if _sha256(output.read_bytes()) != output_hash:
                raise WorkbookError("Published output hash mismatch; stop and quarantine the copy.")
        except FileExistsError as error:
            raise WorkbookError("Output already exists; no overwrite.") from error
        finally:
            exported.close()
            if temporary is not None:
                temporary.unlink(missing_ok=True)
        return {"output": str(output), "source_sha256": _sha256(original),
                "output_sha256": output_hash, "model_row_count": len(rows),
                "source_rows": rows, "columns": columns,
                "assessment": "Allowlisted review copy only; human review and sharing authorization still required."}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    for name in ("inspect", "sharing-export"):
        command = commands.add_parser(name)
        command.add_argument("input", type=Path)
        command.add_argument("--sheet")
        command.add_argument("--first-n", type=int, default=50)
        if name == "sharing-export":
            command.add_argument("--output", required=True, type=Path)
            command.add_argument("--columns", required=True, type=Path, help="JSON config file with columns and trusted_hosts")
    arguments = parser.parse_args(argv)
    try:
        if arguments.command == "inspect":
            result = inspect_workbook(arguments.input, arguments.sheet, arguments.first_n)
        else:
            config = json.loads(arguments.columns.read_text(encoding="utf-8"))
            result = sharing_export(arguments.input, arguments.output, config, arguments.sheet, arguments.first_n)
        print(json.dumps(result, indent=2, default=str))
    except (WorkbookError, OSError, ValueError) as error:
        parser.exit(2, f"aict-workbooks: {error}\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())