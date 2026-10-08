import hashlib
import json
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.hyperlink import Hyperlink

from eai_pmo_tools.aict_workbooks import (
    WorkbookError, inspect_workbook, main, sharing_export,
)


CONFIG = {"columns": ["Model / Sub-component", "AICT ID", "Approved", "Reference"],
          "trusted_hosts": ["surf.service-now.com"]}


def make_source(tmp_path, header_row=3):
    source = tmp_path / "source.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Inventory"
    if header_row > 1:
        sheet.cell(1, 1, "PRIVATE SOURCE TITLE")
    for column, name in enumerate(CONFIG["columns"] + ["Private notes", "Calculation"], 1):
        sheet.cell(header_row, column, name)
    for values in [("Alpha", "AICT1", False, "Review", "SECRET", "=1+1"),
                   ("Beta", "AICT1", True, None, "SECRET", None),
                   ("Gamma", None, False, None, "SECRET", None),
                   ("Delta", "MISSING", False, None, "SECRET", None)]:
        sheet.append(values)
    sheet.cell(header_row + 1, 4).hyperlink = "https://surf.service-now.com/record/123"
    sheet.cell(header_row + 1, 4).hyperlink.tooltip = "SECRET TOOLTIP"
    sheet.cell(header_row + 1, 4).hyperlink.display = "SECRET DISPLAY"
    sheet.cell(header_row + 1, 1).comment = Comment("SECRET COMMENT", "Secret Author")
    sheet.row_dimensions[header_row + 2].hidden = True
    sheet.column_dimensions["C"].hidden = True
    sheet.column_dimensions["E"].hidden = True
    private = book.create_sheet("Private")
    private.append(["SECRET TAB"])
    private.sheet_state = "veryHidden"
    book.properties.creator = "SECRET AUTHOR"
    book.properties.description = "SECRET DESCRIPTION"
    book.save(source)
    book.close()
    return source


def change_cell(source, column, value=None, link=None):
    book = load_workbook(source)
    cell = book["Inventory"].cell(4, column)
    if value is not None:
        cell.value = value
    if link is not None:
        cell.hyperlink = link
    book.save(source)
    book.close()


@pytest.mark.parametrize("header_row", [1, 3, 11])
def test_detects_header_and_independent_missing_duplicate_rows(tmp_path, header_row):
    source = make_source(tmp_path, header_row)
    result = inspect_workbook(source)
    assert result["header_row"] == header_row
    assert result["model_row_count"] == result["retained_model_rows"] == 4
    assert result["duplicate_ids"] == {"AICT1": [header_row + 1, header_row + 2]}
    assert result["missing_id_rows"] == [header_row + 3, header_row + 4]
    assert result["hidden_model_rows"] == 1
    assert result["visible_model_rows"] == 3


def test_export_is_clean_preserves_rows_values_false_and_source_hash(tmp_path):
    source = make_source(tmp_path)
    original = source.read_bytes()
    output = tmp_path / "review.xlsx"
    inspected = inspect_workbook(source)
    result = sharing_export(source, output, CONFIG)
    assert source.read_bytes() == original
    assert inspected["source_sha256"] == result["source_sha256"] == hashlib.sha256(original).hexdigest()
    assert result["output_sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert result["source_rows"] == [4, 5, 6, 7]
    book = load_workbook(output)
    assert book.sheetnames == ["Review"]
    sheet = book.active
    assert sheet.max_row == 5 and sheet.max_column == 4
    assert sheet["C2"].value is False
    assert sheet["B2"].value == sheet["B3"].value == "AICT1"
    assert sheet["B4"].value is None and sheet["B5"].value == "MISSING"
    assert sheet["D2"].hyperlink.target == "https://surf.service-now.com/record/123"
    assert sheet["D2"].hyperlink.tooltip is None and sheet["D2"].hyperlink.display is None
    assert all(cell.comment is None and cell.data_type != "f" for row in sheet for cell in row)
    assert not any(dim.hidden for dim in sheet.row_dimensions.values())
    assert not any(dim.hidden for dim in sheet.column_dimensions.values())
    assert not book.properties.creator and not book.properties.description
    book.close()
    with ZipFile(output) as archive:
        assert b"SECRET" not in b"".join(archive.read(name) for name in archive.namelist())
    assert not list(tmp_path.glob(".aict-review-*"))


def test_side_output_required_and_existing_output_never_overwritten(tmp_path):
    source = make_source(tmp_path)
    original = source.read_bytes()
    with pytest.raises(WorkbookError, match="no overwrite"):
        sharing_export(source, source, CONFIG)
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"KEEP")
    with pytest.raises(WorkbookError, match="no overwrite"):
        sharing_export(source, output, CONFIG)
    assert output.read_bytes() == b"KEEP" and source.read_bytes() == original


def test_symlink_output_rejected(tmp_path):
    source = make_source(tmp_path)
    output = tmp_path / "review.xlsx"
    output.symlink_to(tmp_path / "not-yet-created.xlsx")
    with pytest.raises(WorkbookError, match="no overwrite"):
        sharing_export(source, output, CONFIG)


def test_formula_selected_rejected_without_side_output(tmp_path):
    source = make_source(tmp_path)
    change_cell(source, 3, value="=FALSE()")
    with pytest.raises(WorkbookError, match="formula"):
        sharing_export(source, tmp_path / "review.xlsx", CONFIG)
    assert not (tmp_path / "review.xlsx").exists()


@pytest.mark.parametrize("link", ["javascript:alert(1)", "http://surf.service-now.com/path",
    "https://evil.example/path", "https://surf.service-now.com.evil.example/path",
    "https://user:password@surf.service-now.com/path", "https://surf.service-now.com/path?token=secret",
    "https://surf.service-now.com/path#secret", "https://surf.service-now.com:8443/path",
    "https://surf.service-now.com/path%3Ftoken=secret", "https://surf.service-now.com/path%0Asecret"])
def test_bad_link_rejected(tmp_path, link):
    source = make_source(tmp_path)
    change_cell(source, 4, link=link)
    with pytest.raises(WorkbookError, match="hyperlink|Hyperlink"):
        sharing_export(source, tmp_path / "review.xlsx", CONFIG)
    assert not (tmp_path / "review.xlsx").exists()


def test_internal_link_rejected(tmp_path):
    source = make_source(tmp_path)
    change_cell(source, 4, link=Hyperlink(ref="D4", location="Private!A1"))
    with pytest.raises(WorkbookError, match="Internal"):
        sharing_export(source, tmp_path / "review.xlsx", CONFIG)


@pytest.mark.parametrize("payload, message", [(b"not a zip", "Invalid ZIP"),
    (bytes.fromhex("D0CF11E0A1B11AE1") + b"encrypted", "Encrypted")])
def test_unreadable_input_clear_error_unchanged(tmp_path, payload, message):
    source = tmp_path / "source.xlsx"
    source.write_bytes(payload)
    with pytest.raises(WorkbookError, match=message):
        inspect_workbook(source)
    assert source.read_bytes() == payload


def test_invalid_ooxml_package(tmp_path):
    source = tmp_path / "source.xlsx"
    with ZipFile(source, "w") as archive:
        archive.writestr("hello.txt", "not a workbook")
    with pytest.raises(WorkbookError, match="required workbook parts"):
        inspect_workbook(source)


def test_corrupt_xml_stops(tmp_path):
    source = tmp_path / "source.xlsx"
    with ZipFile(source, "w") as archive:
        archive.writestr("[Content_Types].xml", "broken xml")
        archive.writestr("xl/workbook.xml", "broken xml")
    with pytest.raises(WorkbookError, match="corrupt XLSX"):
        inspect_workbook(source)


@pytest.mark.parametrize("config", [{}, {"columns": [], "trusted_hosts": []},
    {"columns": ["AICT ID", "AICT ID"], "trusted_hosts": []},
    {"columns": ["Unknown"], "trusted_hosts": []},
    {"columns": ["AICT ID"], "trusted_hosts": ["*.service-now.com"]}])
def test_strict_config(tmp_path, config):
    source = make_source(tmp_path)
    with pytest.raises(WorkbookError):
        sharing_export(source, tmp_path / "review.xlsx", config)


def test_multiple_inventory_sheets_require_explicit_sheet(tmp_path):
    source = make_source(tmp_path)
    book = load_workbook(source)
    book.copy_worksheet(book["Inventory"])
    book.save(source)
    book.close()
    with pytest.raises(WorkbookError, match="Expected one"):
        inspect_workbook(source)
    assert inspect_workbook(source, "Inventory")["model_row_count"] == 4


def test_duplicate_named_headers_rejected(tmp_path):
    source = make_source(tmp_path)
    book = load_workbook(source)
    book["Inventory"]["G3"] = "AICT ID"
    book.save(source)
    book.close()
    with pytest.raises(WorkbookError, match="duplicate named headers"):
        inspect_workbook(source)


def test_protection_reported_but_export_not_bypassed(tmp_path):
    source = make_source(tmp_path)
    book = load_workbook(source)
    book["Inventory"].protection.sheet = True
    book.save(source)
    book.close()
    assert inspect_workbook(source)["sheets"][0]["protected"] is True
    with pytest.raises(WorkbookError, match="Protected"):
        sharing_export(source, tmp_path / "review.xlsx", CONFIG)


def test_first_n_bound_and_missing_sheet(tmp_path):
    source = make_source(tmp_path, 11)
    with pytest.raises(WorkbookError):
        inspect_workbook(source, first_n=10)
    with pytest.raises(WorkbookError):
        inspect_workbook(source, first_n=0)
    with pytest.raises(WorkbookError, match="does not exist"):
        inspect_workbook(source, sheet="Absent")


def test_publish_race_does_not_overwrite_and_cleans_temp(tmp_path):
    source = make_source(tmp_path)
    output = tmp_path / "review.xlsx"

    def concurrent_publish(temporary, destination):
        Path(destination).write_bytes(b"OTHER WRITER")
        raise FileExistsError()

    with patch("eai_pmo_tools.aict_workbooks.os.link", side_effect=concurrent_publish):
        with pytest.raises(WorkbookError, match="no overwrite"):
            sharing_export(source, output, CONFIG)
    assert output.read_bytes() == b"OTHER WRITER"
    assert not list(tmp_path.glob(".aict-review-*"))


def test_cli_inspect_and_export(tmp_path, capsys):
    source = make_source(tmp_path)
    config_path = tmp_path / "columns.json"
    config_path.write_text(json.dumps(CONFIG))
    assert main(["inspect", str(source), "--sheet", "Inventory"]) == 0
    assert json.loads(capsys.readouterr().out)["model_row_count"] == 4
    assert main(["sharing-export", str(source), "--columns", str(config_path),
                 "--output", str(tmp_path / "review.xlsx")]) == 0
    assert json.loads(capsys.readouterr().out)["model_row_count"] == 4


def test_empty_allowlisted_values_still_retain_every_model_row(tmp_path):
    source = make_source(tmp_path)
    output = tmp_path / "review.xlsx"
    config = {"columns": ["Calculation"], "trusted_hosts": []}
    change_cell(source, 6, value="text")
    book = load_workbook(source)
    book["Inventory"]["F4"] = None
    book.save(source)
    book.close()
    result = sharing_export(source, output, config)
    book = load_workbook(output)
    assert result["model_row_count"] == 4 and book.active.max_row == 5
    assert [book.active.cell(row, 1).value for row in range(2, 6)] == [None] * 4
    book.close()


def test_literal_formula_looking_text_remains_text(tmp_path):
    source = make_source(tmp_path)
    book = load_workbook(source)
    cell = book["Inventory"]["A4"]
    cell.value = "=not_a_formula"
    cell.data_type = "s"
    book.save(source)
    book.close()
    output = tmp_path / "review.xlsx"
    sharing_export(source, output, CONFIG)
    book = load_workbook(output)
    assert book.active["A2"].value == "=not_a_formula"
    assert book.active["A2"].data_type == "s"
    book.close()


def test_source_change_before_publication_stops_and_cleans_temp(tmp_path):
    source = make_source(tmp_path)
    output = tmp_path / "review.xlsx"

    def changed_source(*arguments):
        raise WorkbookError("Source changed during processing")

    with patch("eai_pmo_tools.aict_workbooks._assert_unchanged", side_effect=changed_source):
        with pytest.raises(WorkbookError, match="Source changed"):
            sharing_export(source, output, CONFIG)
    assert not output.exists() and not list(tmp_path.glob(".aict-review-*"))


def test_unselected_unsafe_links_are_not_copied(tmp_path):
    source = make_source(tmp_path)
    change_cell(source, 5, link="javascript:alert(1)")
    output = tmp_path / "review.xlsx"
    sharing_export(source, output, CONFIG)
    with ZipFile(output) as archive:
        assert b"javascript" not in b"".join(archive.read(name) for name in archive.namelist())


def test_cli_invalid_config_exits_cleanly(tmp_path, capsys):
    source = make_source(tmp_path)
    config_path = tmp_path / "columns.json"
    config_path.write_text("invalid json")
    with pytest.raises(SystemExit) as error:
        main(["sharing-export", str(source), "--columns", str(config_path),
              "--output", str(tmp_path / "review.xlsx")])
    assert error.value.code == 2
    assert "aict-workbooks:" in capsys.readouterr().err